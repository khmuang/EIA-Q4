import os
import sys
import json
import datetime
import openpyxl

# --- CONFIGURATION (OneDrive Integration for Q4) ---
ONEDRIVE_Q4_ROOT = r"D:\Users\Djmanny\OneDrive - Central Group\RIS Endpoint support - 2026\Q4"
Q4_SCOPE_DOC = os.path.join(ONEDRIVE_Q4_ROOT, "EIA Document Q4.xlsx")

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(BASE_DIR, 'data.js')

TOPICS_CONFIG = {
    "1.1 IT Asset Management.xlsx": {"id": "1.1", "subfolder": "1.1 IT Asset Management", "status_col": "Asset update status Y/N", "team_col": "Groups"},
    "1.2 Install GLPI agent.xlsx": {"id": "1.2", "subfolder": "1.2 Install GLPI agent", "status_col": "GLPI setup status Y/N", "team_col": "Serviced By"},
    "2. Update OS.xlsx": {"id": "2", "subfolder": "2. Update OS", "status_col": "OS update status Y/N", "team_col": "Serviced By"},
    "3. Device Require Patch Update.xlsx": {"id": "3", "subfolder": "3. Patch Update", "status_col": "Patch status Y/N", "team_col": "Serviced By"},
    "4. Antivirus Installation.xlsx": {"id": "4", "subfolder": "4. Antivirus Installation", "status_col": "AV update Y/N", "team_col": "Serviced By"},
    "5. Built-in Firewall Enablement.xlsx": {"id": "5", "subfolder": "5. Built-in Firewall Enablement", "status_col": "Firewall update Y/N", "team_col": "Serviced By"},
    "6. Client join domain.xlsx": {"id": "6", "subfolder": "6. Client join domain", "status_col": "Join domain update Y/N", "team_col": "Serviced By"},
    "7. Privileged User management.xlsx": {"id": "7", "subfolder": "7. Privileged User management", "status_col": "Std admin update Y/N", "team_col": "Serviced By"},
    "8. Document Request.xlsx": {"id": "8", "subfolder": "8. Document Request", "status_col": "Document request update Y/N", "team_col": "Serviced By"}
}

def load_baseline_from_q4_document():
    """Extract baseline targets when granular per-machine topic sheets are not yet available."""
    print(f"[INFO] Parsing Baseline Scope from: {Q4_SCOPE_DOC}")
    if not os.path.exists(Q4_SCOPE_DOC):
        print(f"[ERROR] Baseline scope document not found: {Q4_SCOPE_DOC}")
        return None

    wb = openpyxl.load_workbook(Q4_SCOPE_DOC, data_only=True)
    sheet = wb.active

    scope_totals = {}
    for r in range(4, sheet.max_row + 1):
        no_val = str(sheet.cell(row=r, column=2).value or "").strip()
        total_val = sheet.cell(row=r, column=4).value
        if no_val and isinstance(total_val, (int, float)):
            # Normalize topic keys
            topic_key = no_val.replace(".1", "").replace(".2", "") if no_val not in ["1.1", "1.2"] else no_val
            if no_val == "1.1": topic_key = "1.1"
            elif no_val == "1.2": topic_key = "1.2"
            elif no_val.startswith("2"): topic_key = "2"
            elif no_val.startswith("3"): topic_key = "3"
            elif no_val.startswith("4"): topic_key = "4"
            elif no_val.startswith("5"): topic_key = "5"
            elif no_val.startswith("6"): topic_key = "6"
            elif no_val.startswith("7"): topic_key = "7"
            elif no_val.startswith("8"): topic_key = "8"
            
            scope_totals[topic_key] = int(total_val)

    wb.close()

    # Distribute baseline totals across standard support teams (HO: 45%, Branch: 40%, DC: 15%)
    # to provide immediate, realistic executive preview until granular files arrive
    multi_matrix = {"HO": {}, "Branch": {}, "DC": {}}
    ratios = {"HO": 0.45, "Branch": 0.40, "DC": 0.15}

    for topic_id, total in scope_totals.items():
        allocated = 0
        teams = list(ratios.keys())
        for idx, team in enumerate(teams):
            if idx == len(teams) - 1:
                t_count = total - allocated
            else:
                t_count = int(total * ratios[team])
                allocated += t_count

            multi_matrix[team][topic_id] = {
                "Q4": {
                    "total": t_count,
                    "success": 0, # Initial progress before resolution
                    "bu_breakdown": {
                        "Corporate": {
                            "total": t_count,
                            "success": 0
                        }
                    }
                }
            }

    return multi_matrix

def sync():
    print(f"==================================================")
    print(f"       EIA Q4 DATA SYNC & EXTRACTION ENGINE       ")
    print(f"==================================================")
    print(f"[PATH] OneDrive Q4 Source: {ONEDRIVE_Q4_ROOT}")
    
    # Check if granular topic files exist in OneDrive Q4
    has_granular = False
    for filename, mapping in TOPICS_CONFIG.items():
        file_path = os.path.join(ONEDRIVE_Q4_ROOT, mapping['subfolder'], filename)
        alt_path = os.path.join(ONEDRIVE_Q4_ROOT, filename)
        if os.path.exists(file_path) or os.path.exists(alt_path):
            has_granular = True
            break

    multi_matrix = None
    if has_granular:
        print("[MODE] Found granular topic Excel files. Processing per-machine records...")
        import pandas as pd
        multi_matrix = {}
        for filename, mapping in TOPICS_CONFIG.items():
            topic_id = mapping['id']
            # Resolve subfolder
            subfolder_path = os.path.join(ONEDRIVE_Q4_ROOT, mapping['subfolder'])
            if not os.path.exists(subfolder_path):
                # Try finding subfolder by topic ID prefix
                for d in os.listdir(ONEDRIVE_Q4_ROOT):
                    if os.path.isdir(os.path.join(ONEDRIVE_Q4_ROOT, d)):
                        if d.startswith(topic_id + " ") or d.startswith(topic_id + "."):
                            subfolder_path = os.path.join(ONEDRIVE_Q4_ROOT, d)
                            break

            file_path = os.path.join(subfolder_path, filename)
            if not os.path.exists(file_path) and os.path.exists(subfolder_path):
                # Look for matching excel file in subfolder
                cands = [f for f in os.listdir(subfolder_path) if f.endswith('.xlsx') and not f.startswith('~$') and 'backup' not in f.lower()]
                if cands:
                    file_path = os.path.join(subfolder_path, cands[0])

            if not os.path.exists(file_path):
                file_path = os.path.join(ONEDRIVE_Q4_ROOT, filename)

            if not os.path.exists(file_path):
                print(f"[SKIP] Topic file not found yet: {filename}")
                continue

            try:
                df = pd.read_excel(file_path, header=0)
                status_col = mapping['status_col']
                team_col = mapping['team_col']
                if status_col not in df.columns:
                    yn_cols = [c for c in df.columns if 'Y/N' in str(c).upper() or 'STATUS' in str(c).upper()]
                    if yn_cols: status_col = yn_cols[0]

                if team_col not in df.columns:
                    team_cands = [c for c in df.columns if 'SERVICE' in str(c).upper() or 'GROUP' in str(c).upper() or 'TEAM' in str(c).upper()]
                    if team_cands: team_col = team_cands[0]

                bu_col = None
                if 'BU' in df.columns: bu_col = 'BU'
                elif 'Plugins - BU, Company - BU' in df.columns: bu_col = 'Plugins - BU, Company - BU'
                else:
                    found_bu = [c for c in df.columns if 'BU' in str(c).upper() and 'BUILD' not in str(c).upper()]
                    if found_bu: bu_col = found_bu[0]

                if status_col and team_col:
                    df['team_clean'] = df[team_col].fillna('Unknown').astype(str).str.strip()
                    df['y_n'] = df[status_col].fillna('N').apply(lambda x: 'Y' if str(x).strip().upper() == 'Y' else 'N')
                    if 'EIA Phase' in df.columns:
                        df['phase_clean'] = df['EIA Phase'].fillna('Q4').astype(str).str.strip().str.split(',').str[0]
                    else:
                        df['phase_clean'] = 'Q4'

                    if bu_col:
                        df['bu_clean'] = df[bu_col].fillna('Unknown').astype(str).str.strip()
                    else:
                        df['bu_clean'] = 'Unknown'

                    for (team, phase), group in df.groupby(['team_clean', 'phase_clean']):
                        if team not in multi_matrix: multi_matrix[team] = {}
                        if topic_id not in multi_matrix[team]: multi_matrix[team][topic_id] = {}

                        bu_breakdown = {}
                        for bu, bu_group in group.groupby('bu_clean'):
                            bu_breakdown[bu] = {
                                "total": int(len(bu_group)),
                                "success": int((bu_group['y_n'] == 'Y').sum())
                            }

                        multi_matrix[team][topic_id][phase] = {
                            "total": int(len(group)),
                            "success": int((group['y_n'] == 'Y').sum()),
                            "bu_breakdown": bu_breakdown
                        }
            except Exception as e:
                print(f"[ERROR] processing {filename}: {e}")
    else:
        print("[MODE] No granular topic folders yet. Initializing Baseline Scope from EIA Document Q4.xlsx...")
        multi_matrix = load_baseline_from_q4_document()

    if not multi_matrix:
        print("[ERROR] Failed to generate matrix. Aborting.")
        sys.exit(1)

    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    js_content = f"// Automatically generated for EIA Q4 at {now}\nconst DASHBOARD_DATA = {json.dumps(multi_matrix, indent=4)};\n"
    js_content += f"const LAST_UPDATED = '{now}';\n"
    js_content += f"const ACTIVE_QUARTER = 'Q4';\n"

    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        f.write(js_content)

    print(f"[SUCCESS] EIA Q4 data.js generated successfully!")
    print(f"[OUTPUT] {OUTPUT_FILE}")

if __name__ == "__main__":
    sync()
